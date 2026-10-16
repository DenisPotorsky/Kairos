"""Парсер расписания залов Консерватории (УГК) — с фильтром Стэнвей."""

import re
from datetime import datetime, timedelta
from pathlib import Path
from loguru import logger

from kairos.models import Event, EventType, Priority
from kairos.parsers.base import BaseParser


class ExcelUgkParser(BaseParser):
    ORG_NAME = "Консерватория"
    # ИСПРАВЛЕНО: добавили "стэнвей" (через э)
    ALLOWED_KEYWORDS = [
        "концерт",
        "настройка",
        "экзамен",
        "зачет",
        "запись",
        "мастер-класс",
        "мастер класс",
        "стенвей",
        "стейнвей",
        "стэнвей"  # <-- ДОБАВИЛИ ЭТОТ ВАРИАНТ
    ]

    def parse(self, file_path: Path) -> list[Event]:
        if file_path.suffix.lower() not in [".xls", ".xlsx"]:
            return []

        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("❌ Ошибка: pip install openpyxl")
            return []

        events = []
        try:
            wb = load_workbook(filename=file_path, data_only=True)

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_events = self._parse_sheet(ws, file_path.name)
                events.extend(sheet_events)

            wb.close()
            logger.info(f"✅ УГК: найдено {len(events)} событий")

        except Exception as e:
            logger.error(f"❌ Ошибка чтения файла УГК: {e}")

        return events

    def _parse_sheet(self, ws, source: str) -> list[Event]:
        events = []
        current_date = None
        current_hall = "Зал УГК"

        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            full_row_text = " ".join(cells).lower()

            # 1. ПРОВЕРКА НА СМЕНУ ЗАЛА
            if "малый зал" in full_row_text:
                current_hall = "Малый зал УГК"
                continue

            if "большой зал" in full_row_text:
                current_hall = "Большой зал УГК"
                continue

            # 2. ПОИСК ДАТЫ
            date_str = cells[0]
            date_match = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", date_str)
            if date_match:
                day = int(date_match.group(1))
                month = int(date_match.group(2))
                year = int(date_match.group(3))
                if year < 100:
                    year += 2000
                try:
                    current_date = datetime(year, month, day)
                except ValueError:
                    continue

            if not current_date:
                continue

            # 3. ПАРСИНГ СОБЫТИЯ
            time_str = cells[1] if len(cells) > 1 else ""
            title_str = cells[2] if len(cells) > 2 else ""

            if not time_str or not title_str:
                continue
            if "дата" in title_str.lower() or "время" in title_str.lower():
                continue

            # 4. ФИЛЬТР
            lower_title = title_str.lower()
            if not any(kw in lower_title for kw in self.ALLOWED_KEYWORDS):
                continue

            # 5. СОЗДАНИЕ СОБЫТИЯ
            event = self._create_event(current_date, time_str, title_str, current_hall, source)
            if event:
                events.append(event)

        return events

    def _create_event(self, base_date: datetime, time_str: str, title: str, hall: str, source: str) -> Event | None:
        clean_time = time_str.replace(".", ":").replace(" ", "")

        match = re.search(r"(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})", clean_time)
        if match:
            start_dt = base_date.replace(hour=int(match.group(1)), minute=int(match.group(2)))
            end_dt = base_date.replace(hour=int(match.group(3)), minute=int(match.group(4)))
        else:
            single_match = re.search(r"(\d{1,2}):(\d{2})", clean_time)
            if single_match:
                h, m = int(single_match.group(1)), int(single_match.group(2))
                start_dt = base_date.replace(hour=h, minute=m)
                end_dt = start_dt + timedelta(hours=2)
            else:
                return None

        lower_title = title.lower()
        event_type = EventType.REHEARSAL

        if "концерт" in lower_title:
            event_type = EventType.CONCERT
        elif "настройка" in lower_title:
            event_type = EventType.TECHNICAL
        elif "экзамен" in lower_title or "зачет" in lower_title or "экз" in lower_title:
            event_type = EventType.MEETING
        elif "запись" in lower_title:
            event_type = EventType.SERVICE
        elif "мастер" in lower_title:
            event_type = EventType.MEETING
        elif "стенвей" in lower_title or "стейнвей" in lower_title or "стэнвей" in lower_title:
            event_type = EventType.TECHNICAL

        return Event(
            title=f"[{self.ORG_NAME}] {title}",
            start=start_dt,
            end=end_dt,
            event_type=event_type,
            priority=Priority.P1,
            source_file=source,
            location=hall,
            tags=()
        )
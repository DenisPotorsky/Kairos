"""Парсер сезонного плана Филармонии (DOCX) — полная версия."""

import re
from datetime import datetime, timedelta
from pathlib import Path
from loguru import logger

from docx import Document

from kairos.models import Event, EventType, Priority
from kairos.parsers.base import BaseParser


class DocxSeasonParser(BaseParser):
    ORG_NAME = "Сезон Хора"

    def parse(self, file_path: Path) -> list[Event]:
        self.validate_file(file_path, ".docx")
        events = []

        try:
            doc = Document(file_path)

            # 1. Собираем весь текст из ТАБЛИЦ (основной источник)
            full_text = ""
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text += cell.text + "\n"

            # 2. Добавляем текст из параграфов (на всякий случай)
            for para in doc.paragraphs:
                full_text += para.text + "\n"

            # Разбиваем на строки для обработки
            lines = full_text.split('\n')

            for line in lines:
                line = line.strip()
                if not line:
                    continue

                # Ищем дату в формате ДД.ММ
                # Примеры строк:
                # "11 .09,  пт, 19 :00,  Аб. 5 Бриттен..."
                # "28.10, ср. Зарядье Гаврилин..."

                # Регулярка: День.Месяц, (опционально день недели), (опционально время), Описание
                match = re.search(r'(\d{1,2})\s*\.\s*(\d{1,2}),?\s*[а-яёA-Za-z]*\.?\s*,?\s*(?:(\d{1,2})\s*:\s*(\d{2}))?\s*,?\s*(.*)', line)

                if match:
                    day = int(match.group(1))
                    month = int(match.group(2))
                    hour = match.group(3)
                    minute = match.group(4)
                    description = match.group(5).strip()

                    # Убираем мусор в начале описания (например "Аб. 5")
                    # Но оставляем само название концерта
                    description = re.sub(r'^Аб\.\s*\d+\s*', '', description)

                    if not description:
                        continue

                    try:
                        # Определяем год (Сезон 2026-2027)
                        year = 2026
                        if month < 9: # Янв-Авг -> 2027
                            year = 2027

                        base_date = datetime(year, month, day)

                        start_dt = None
                        end_dt = None

                        if hour and minute:
                            start_dt = base_date.replace(hour=int(hour), minute=int(minute))
                            end_dt = start_dt + timedelta(hours=2)
                        else:
                            # Без времени -> весь день
                            start_dt = base_date.replace(hour=0, minute=0)
                            end_dt = base_date.replace(hour=23, minute=59)

                        # Формируем событие
                        # В title пишем полное описание, чтобы было видно в календаре
                        events.append(Event(
                            title=f"[{self.ORG_NAME}] {description}",
                            start=start_dt,
                            end=end_dt,
                            event_type=EventType.CONCERT,
                            priority=Priority.P0,
                            source_file=file_path.name,
                            location="Филармония"
                        ))

                    except ValueError:
                        continue

        except Exception as e:
            logger.error(f"❌ Ошибка парсинга сезона {file_path.name}: {e}")

        logger.info(f"✅ Сезон Хора: {len(events)} событий")
        return events
"""Парсер плана залов филармонии — читает сконвертированный Excel."""

import re
from datetime import datetime, timedelta
from pathlib import Path
from loguru import logger

from openpyxl import load_workbook

from kairos.models import Event, EventType, Priority
from kairos.parsers.base import BaseParser


class PdfHallScheduleParser(BaseParser):
    ORG_NAME = "Филармония"
    ALLOWED_KEYWORDS = ["петренко", "рояль", "спинет"]

    def parse(self, file_path: Path) -> list[Event]:
        # file_path здесь приходит .pdf, но мы ищем .xlsx рядом
        excel_path = file_path.with_suffix('.xlsx')

        if not excel_path.exists():
            logger.warning(f"⚠️ Excel файл {excel_path.name} не найден. Запустите конвертацию.")
            return []

        events = []
        try:
            wb = load_workbook(excel_path, read_only=True, data_only=True)
            ws = wb.active

            for row in ws.iter_rows(min_row=2, values_only=True): # пропускаем заголовок
                if not row or len(row) < 3:
                    continue

                date_str, time_str, desc = str(row[0]), str(row[1]), str(row[2])

                # ФИЛЬТР
                if not any(kw in desc.lower() for kw in self.ALLOWED_KEYWORDS):
                    continue

                # ПАРСИНГ ДАТЫ
                # date_str может быть "21.9" или "28.09"
                m = re.match(r'(\d{1,2})\.(\d{1,2})', date_str)
                if not m:
                    continue # Пропускаем строки без даты (NO_DATE)

                day, month = int(m.group(1)), int(m.group(2))
                try:
                    base_date = datetime(datetime.now().year, month, day)
                except ValueError:
                    continue

                # ПАРСИНГ ВРЕМЕНИ
                # time_str может быть "10:00-14:00" или "19:00"
                start_dt, end_dt = None, None

                time_range = re.search(r'(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})', time_str)
                if time_range:
                    start_dt = base_date.replace(hour=int(time_range.group(1)), minute=int(time_range.group(2)))
                    end_dt = base_date.replace(hour=int(time_range.group(3)), minute=int(time_range.group(4)))
                else:
                    time_single = re.search(r'(\d{1,2}):(\d{2})', time_str)
                    if time_single:
                        start_dt = base_date.replace(hour=int(time_single.group(1)), minute=int(time_single.group(2)))
                        end_dt = start_dt + timedelta(hours=2)

                if not start_dt:
                    # Если времени нет, ставим заглушку на 12:00
                    start_dt = base_date.replace(hour=12, minute=0)
                    end_dt = start_dt + timedelta(hours=1)

                # ТИП СОБЫТИЯ
                etype = EventType.REHEARSAL
                if 'концерт' in desc.lower():
                    etype = EventType.CONCERT
                elif 'настройка' in desc.lower() or 'спинет' in desc.lower() or 'рояль' in desc.lower():
                    etype = EventType.TECHNICAL

                events.append(Event(
                    title=f"[{self.ORG_NAME}] {desc}",
                    start=start_dt,
                    end=end_dt,
                    event_type=etype,
                    priority=Priority.P1,
                    source_file=file_path.name,
                    location="Филармония"
                ))

            wb.close()
            logger.info(f"✅ Залы Филармонии (Excel): {len(events)} событий")

        except Exception as e:
            logger.error(f"Ошибка чтения Excel залов: {e}")

        return events